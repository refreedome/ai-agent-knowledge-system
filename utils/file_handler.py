import os
import hashlib
from utils.logger_handler import logger
from langchain_core.documents import Document
from langchain_community.document_loaders import PyPDFLoader, TextLoader


def get_file_md5_hex(filepath: str):     # 获取文件的md5的十六进制字符串

    if not os.path.exists(filepath):
        logger.error(f"[md5计算]文件{filepath}不存在")
        return

    if not os.path.isfile(filepath):
        logger.error(f"[md5计算]路径{filepath}不是文件")
        return

    md5_obj = hashlib.md5()

    chunk_size = 4096       # 4KB分片，避免文件过大爆内存
    try:
        with open(filepath, "rb") as f:     # 必须二进制读取
            while chunk := f.read(chunk_size):
                md5_obj.update(chunk)

            """
            chunk = f.read(chunk_size)
            while chunk:
                
                md5_obj.update(chunk)
                chunk = f.read(chunk_size)
            """
            md5_hex = md5_obj.hexdigest()
            return md5_hex
    except Exception as e:
        logger.error(f"计算文件{filepath}md5失败，{str(e)}")
        return None


def listdir_with_allowed_type(path: str, allowed_types: tuple[str]):        # 返回文件夹内的文件列表（允许的文件后缀）
    files = []

    if not os.path.isdir(path):
        logger.error(f"[listdir_with_allowed_type]{path}不是文件夹")
        return allowed_types

    for f in os.listdir(path):
        if f.endswith(allowed_types):
            files.append(os.path.join(path, f))

    return tuple(files)


def pdf_loader(filepath: str, passwd=None) -> list[Document]:
    return PyPDFLoader(filepath, passwd).load()


def txt_loader(filepath: str) -> list[Document]:
    return TextLoader(filepath, encoding="utf-8").load()
# ---- 新增：docx 加载器 ----
def docx_loader(filepath: str) -> list[Document]:
    """加载 .docx 文件，提取所有段落文本"""
    try:
        from docx import Document as DocxDocument
    except ImportError:
        logger.error("请先安装 python-docx: pip install python-docx")
        return []

    doc = DocxDocument(filepath)
    # 提取所有非空段落
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    text = "\n".join(paragraphs)

    # 也提取表格中的文本
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                text += "\n" + " | ".join(cells)

    if not text.strip():
        logger.warning(f"[docx_loader] {filepath} 中未提取到有效文本")
        return []

    return [Document(page_content=text, metadata={"source": os.path.basename(filepath)})]


# ---- 新增：csv 加载器 ----
def csv_loader(filepath: str) -> list[Document]:
    """加载 .csv 文件，将每一行转为 key:value 文本"""
    import csv

    rows_text = []
    rows_count = 0
    with open(filepath, "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        if not reader.fieldnames:
            logger.warning(f"[csv_loader] {filepath} 无有效列名")
            return []

        for row in reader:
            rows_count += 1
            # 将一行转为 "列名:值" 格式
            line = "，".join([f"{k}:{v}" for k, v in row.items() if v and v.strip()])
            if line:
                rows_text.append(line)

    if not rows_text:
        logger.warning(f"[csv_loader] {filepath} 无有效数据行")
        return []

    text = "\n".join(rows_text)
    doc = Document(
        page_content=text,
        metadata={
            "source": os.path.basename(filepath),
            "rows": rows_count,
            "format": "csv"
        }
    )
    return [doc]


# ---- 新增：图片 OCR 加载器 ----
def image_loader(filepath: str) -> list[Document]:
    """OCR 识别图片中的文字"""
    try:
        from PIL import Image
        import pytesseract
    except ImportError:
        logger.error("请先安装依赖: pip install Pillow pytesseract")
        return []

    # Windows 下可能需要指定 Tesseract 路径，取消下面注释并修改路径
    # import pytesseract
    # pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'

    try:
        img = Image.open(filepath)
        # 尝试中文 + 英文识别
        text = pytesseract.image_to_string(img, lang="chi_sim+eng")
    except Exception as e:
        logger.error(f"[image_loader] OCR 识别失败: {str(e)}")
        # 降级：输出图片基本信息
        text = f"[图片文件: {os.path.basename(filepath)}，OCR 识别失败: {str(e)}]"
        return [Document(page_content=text, metadata={"source": os.path.basename(filepath), "ocr_error": str(e)})]

    if not text.strip():
        logger.warning(f"[image_loader] {filepath} 未识别到文字")
        text = f"[图片文件: {os.path.basename(filepath)}，未识别到文字内容]"

    return [Document(page_content=text, metadata={"source": os.path.basename(filepath)})]


# ---- 新增：文件大小检查 ----
def get_file_size_mb(filepath: str) -> float:
    """获取文件大小（MB）"""
    return os.path.getsize(filepath) / (1024 * 1024)